"""Discover public state results pages and parse explicitly structured data."""

from __future__ import annotations

import csv
import io
import json
import re
from datetime import datetime, timezone
from typing import Any
from urllib.parse import urldefrag, urljoin, urlparse

from bs4 import BeautifulSoup
from openpyxl import load_workbook

from .normalize import HEADER_ALIASES, _county_fips_or_none, _header_key, _integer_or_none, _number_or_none, parse_html_result_tables

DATA_SUFFIXES = (".json", ".csv", ".xlsx")
RESULT_WORDS = ("result", "results", "tally", "canvass", "returns", "unofficial-results", "precinct-results")
NON_HTML_SUFFIXES = (".pdf", ".doc", ".docx", ".zip", ".jpg", ".jpeg", ".png")


def is_target_election_link(link: dict[str, str], election_year: int) -> bool:
    label = f"{link.get('label', '')} {link.get('url', '')}"
    years = [int(year) for year in re.findall(r"(?<!\d)(?:19|20)\d{2}(?!\d)", label)]
    if any(year != election_year for year in years):
        return False
    short_years = [int(year) for year in re.findall(r"(?<!\d)(\d{2})(?:gen|pri|primary|runoff|general)", label, flags=re.I)]
    if any(year != election_year % 100 for year in short_years):
        return False
    return (
        str(election_year) in label
        or election_year % 100 in short_years
        or link.get("electionYearContext") == str(election_year)
    )


def _is_relevant_link(text: str, href: str) -> bool:
    if text.strip().lower().startswith("skip to "):
        return False
    parsed = urlparse(href)
    if parsed.path.lower().endswith(NON_HTML_SUFFIXES):
        return False
    label = f"{text} {parsed.path}".lower().replace("_", "-")
    has_result_term = any(re.search(rf"\b{re.escape(word)}\b", label) for word in RESULT_WORDS)
    return has_result_term or (
        "2026" in label and any(word in label for word in ("election", "vote", "general"))
    )


def _same_host(first_url: str, second_url: str) -> bool:
    return urlparse(first_url).hostname == urlparse(second_url).hostname


def discover_result_links(html: bytes | str, page_url: str) -> list[dict[str, str]]:
    soup = BeautifulSoup(html, "html.parser")
    discovered: dict[str, dict[str, str]] = {}
    page_url_parts = urlparse(page_url)
    for anchor in soup.find_all("a", href=True):
        href, _ = urldefrag(urljoin(page_url, anchor["href"]))
        if urlparse(href).scheme not in {"http", "https"}:
            continue
        href_parts = urlparse(href)
        if (href_parts.scheme, href_parts.netloc, href_parts.path.rstrip("/")) == (
            page_url_parts.scheme, page_url_parts.netloc, page_url_parts.path.rstrip("/")
        ) and not href_parts.query:
            continue
        text = anchor.get_text(" ", strip=True)
        if urlparse(href).path.lower().endswith(DATA_SUFFIXES):
            discovered[href] = {"url": href, "label": text[:200], "kind": "machine-readable-file"}
        elif "county" in f"{text} {href}".lower() and "website" in f"{text} {href}".lower():
            discovered[href] = {"url": href, "label": text[:200], "kind": "official-county-results-directory"}
        elif _is_relevant_link(text, href):
            discovered[href] = {"url": href, "label": text[:200], "kind": "linked-page"}
    return list(discovered.values())


def _normalize_rows(rows: list[dict[str, Any]], *, state: str, source_name: str, source_url: str, captured_at: str, method: str) -> list[dict[str, Any]]:
    normalized: list[dict[str, Any]] = []
    for row in rows:
        keys = {_header_key(str(key)): key for key in row}
        indexes = {
            name: next((keys[alias] for alias in aliases if alias in keys), None)
            for name, aliases in HEADER_ALIASES.items()
        }
        if indexes["candidate"] is None or indexes["votes"] is None or indexes["race"] is None:
            continue
        candidate = str(row.get(indexes["candidate"]) or "").strip()
        votes = _integer_or_none(str(row.get(indexes["votes"]) or ""))
        race = str(row.get(indexes["race"]) or "").strip()
        if not candidate or votes is None or not race:
            continue
        county = str(row.get(indexes["county"]) or "").strip() if indexes["county"] else ""
        county_fips = str(row.get(indexes["county_fips"]) or "").strip() if indexes["county_fips"] else ""
        district = str(row.get(indexes["district"]) or "").strip() if indexes["district"] else ""
        precinct = str(row.get(indexes["precinct"]) or "").strip() if indexes["precinct"] else ""
        precinct_id = str(row.get(indexes["precinct_id"]) or "").strip() if indexes["precinct_id"] else ""
        precinct_fips = str(row.get(indexes["precinct_fips"]) or "").strip() if indexes["precinct_fips"] else ""
        result = {
            "state": state,
            "office": race,
            "race": race,
            "candidate": candidate,
            "party": str(row.get(indexes["party"]) or "").strip() if indexes["party"] else "",
            "votes": votes,
            "candidatePct": _number_or_none(str(row.get(indexes["candidate_pct"]) or "")) if indexes["candidate_pct"] else None,
            "precinctsReportingPct": _number_or_none(str(row.get(indexes["reporting"]) or "")) if indexes["reporting"] else None,
            "raceLeaderParty": "",
            "winnerDeclared": False,
            "source": source_name,
            "sourceUrl": source_url,
            "sourceAsOf": None,
            "capturedAt": captured_at,
            "feedLive": None,
            "feedStale": None,
            "method": method,
            "verificationRequired": True,
        }
        if county:
            result["county"] = county
        if district:
            result["district"] = district
        if precinct:
            result["precinct"] = precinct
        if precinct_id:
            result["precinctId"] = precinct_id
        if precinct_fips:
            result["precinctFips"] = precinct_fips
        for field, normalized_field in (
            ("precincts_reported", "precinctsReported"),
            ("precincts_total", "precinctsTotal"),
            ("votes_outstanding", "votesOutstanding"),
            ("ballots_outstanding", "ballotsOutstanding"),
        ):
            index = indexes[field]
            value = _integer_or_none(str(row.get(index) or "")) if index else None
            if value is not None and value >= 0:
                result[normalized_field] = value
        if county_fips:
            normalized_fips = _county_fips_or_none(county_fips, state)
            if normalized_fips:
                result["countyFips"] = normalized_fips
        normalized.append(result)
    return normalized


def _matrix_rows(matrix: list[list[Any]]) -> list[dict[str, Any]]:
    for header_index, raw_headers in enumerate(matrix[:12]):
        headers = [_header_key(str(cell or "")) for cell in raw_headers]
        candidate_match = any(header in HEADER_ALIASES["candidate"] for header in headers)
        votes_match = any(header in HEADER_ALIASES["votes"] for header in headers)
        if candidate_match and votes_match:
            rows = []
            for values in matrix[header_index + 1:]:
                padded = list(values[:len(headers)]) + [None] * max(0, len(headers) - len(values))
                rows.append(dict(zip(headers, padded)))
            return rows
    return []


def parse_machine_file(content: bytes, url: str, *, state: str, source_name: str, captured_at: str) -> list[dict[str, Any]]:
    path = urlparse(url).path.lower()
    if path.endswith(".csv"):
        text = content.decode("utf-8-sig", errors="replace")
        try:
            dialect = csv.Sniffer().sniff(text[:4096])
        except csv.Error:
            dialect = csv.excel
        rows = list(csv.DictReader(io.StringIO(text), dialect=dialect))
        return _normalize_rows(rows, state=state, source_name=source_name, source_url=url, captured_at=captured_at, method="csv-header-match")

    if path.endswith(".xlsx"):
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        normalized = []
        for worksheet in workbook.worksheets:
            rows = _matrix_rows([list(row) for row in worksheet.iter_rows(values_only=True)])
            normalized.extend(_normalize_rows(rows, state=state, source_name=source_name, source_url=url, captured_at=captured_at, method="xlsx-header-match"))
        workbook.close()
        return normalized

    if path.endswith(".json"):
        payload = json.loads(content.decode("utf-8-sig"))
        if isinstance(payload, dict) and isinstance(payload.get("results"), list):
            rows = payload["results"]
        elif isinstance(payload, list):
            rows = payload
        else:
            rows = []
        if not all(isinstance(row, dict) for row in rows):
            return []
        return _normalize_rows(rows, state=state, source_name=source_name, source_url=url, captured_at=captured_at, method="json-header-match")

    return []


def inspect_state_page(page: dict[str, Any], html: bytes, *, captured_at: str) -> tuple[list[dict[str, Any]], list[dict[str, str]]]:
    state = page["state"]
    name = page["name"]
    url = page["results_page"]
    results = parse_html_result_tables(html, state=state, source_name=f"{name} official results page", source_url=url, captured_at=captured_at)
    links = discover_result_links(html, url)
    for link in links:
        if link["kind"] == "machine-readable-file" and not _same_host(url, link["url"]):
            link["kind"] = "external-machine-readable-file-review-required"
    return results, links


def now_utc() -> str:
    return datetime.now(timezone.utc).isoformat(timespec="seconds")
